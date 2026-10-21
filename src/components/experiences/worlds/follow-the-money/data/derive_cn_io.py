"""Derive the China supply-chain splits used by the follow-the-money world.

Input: the official NBS 2023 competitive input-output table (211 product
sectors, 万元, producers' prices), downloaded from the NBS data portal:
https://data.stats.gov.cn/dg/website/publicrelease.pblib/pblibedit/upload/file/xlsx/894bad73a17f4b5b88102f6310584a31
SHA-256 of the file used: 71768c30b8b09aca518cc1d0a316ccdcc525fbe620586d620583c1b986692e95

Usage:
  python3 -I derive_cn_io.py <io2023_competitive.xlsx> <out.json> [--list-sectors]

Method (all money in yuan per 100 yuan of spending):
  * A = Z / x (technical coefficients). The table is competitive-import:
    imports sit inside every use cell. Each product's import share of
    domestic use s_i = IM_i / (TIU_i + TFU_i - EX_i) is assumed equal for all
    its users (the standard proportionality assumption).
  * Domestic coefficients A_d = diag(1 - s) A; imported inputs a_m = s' A.
  * A purchase vector y is split into direct imports s∘y and domestic
    demand y_d = (1 - s)∘y. Purchases that happen inside China (a meal, rent,
    a doctor visit, tuition, electricity) set s = 0 for the purchased item.
  * Goods bought by households carry the average urban trade margin: the
    household purchases of retail and wholesale services divided by
    household purchases of goods (sectors 01-43).
  * Round k of the supply chain receives q_k = A_d^k y_d. It keeps value
    added (compensation, operating surplus, depreciation, net production
    taxes) and pays imported inputs; the rest, A_d q_k, moves one hop
    upstream. Summed over all rounds, value added plus imports equals 100.
"""
import hashlib
import json
import sys

import numpy as np
import openpyxl

HOPS = 7
EXPECTED_SHA = '71768c30b8b09aca518cc1d0a316ccdcc525fbe620586d620583c1b986692e95'

# Plain-language sector names for every sector that appears in a displayed hop.
SECTOR_NAMES = {
    '01001': ('粮食', 'grain farms'),
    '01002': ('蔬菜水果等种植', 'vegetable and fruit farms'),
    '03004': ('家禽养殖', 'poultry farms'),
    '03005': ('畜牧养殖', 'livestock farms'),
    '04006': ('水产养殖', 'fish farms'),
    '06009': ('煤矿', 'coal mines'),
    '07010': ('原油开采', 'oil fields'),
    '11016': ('采矿服务', 'mining services'),
    '05008': ('农业服务', 'farm services'),
    '13017': ('磨面碾米', 'flour and rice mills'),
    '13018': ('饲料', 'animal feed'),
    '13019': ('食用油', 'cooking oil'),
    '13021': ('肉类加工', 'meat processing'),
    '13024': ('其他农副食品', 'other food processing'),
    '14028': ('其他食品', 'other foods'),
    '15029': ('酒厂', 'distilleries and breweries'),
    '16032': ('烟草', 'tobacco makers'),
    '17033': ('棉纺织', 'cotton textiles'),
    '17034': ('化纤纺织', 'synthetic textiles'),
    '17036': ('印染', 'dyeing and finishing'),
    '28067': ('化纤', 'synthetic fibres'),
    '18040': ('服装厂', 'garment factories'),
    '22046': ('造纸', 'paper mills'),
    '23048': ('印刷包装', 'printing and packaging'),
    '25054': ('炼油厂', 'oil refineries'),
    '26056': ('化工原料', 'basic chemicals'),
    '26057': ('化肥', 'fertiliser'),
    '26059': ('合成材料', 'plastics and resins'),
    '26061': ('专用化学品', 'specialty chemicals'),
    '29070': ('塑料制品', 'plastic products'),
    '30071': ('水泥建材', 'cement and building materials'),
    '31079': ('钢铁', 'steel'),
    '32084': ('有色金属', 'non-ferrous metals'),
    '39129': ('手机整机', 'phone assembly'),
    '39132': ('电子器件', 'chips and electronic devices'),
    '39133': ('电子元件', 'electronic components'),
    '44141': ('电网', 'the power grid'),
    '44139': ('火电厂', 'coal and gas power plants'),
    '44140': ('水电风电等发电', 'hydro, wind and other power'),
    '45143': ('燃气', 'gas supply'),
    '47145': ('房屋建筑', 'building construction'),
    '50154': ('装修', 'fit-out and renovation'),
    '51156': ('批发商', 'wholesalers'),
    '52157': ('零售店', 'retailers'),
    '54161': ('货运物流', 'road freight'),
    '62172': ('餐馆', 'restaurants'),
    '64175': ('互联网服务', 'internet services'),
    '66178': ('银行', 'banks'),
    '68181': ('保险', 'insurance'),
    '70183': ('房地产', 'real estate'),
    '71184': ('租赁', 'leasing'),
    '72185': ('法律咨询广告', 'legal, consulting and advertising'),
    '72186': ('商务服务', 'business services'),
    '80198': ('居民服务', 'personal services'),
    '83201': ('学校', 'schools'),
    '84202': ('医院诊所', 'hospitals and clinics'),
    '27064': ('化学药品', 'pharmaceuticals'),
    '27065': ('中药', 'traditional medicines'),
    '27063': ('原料药', 'drug ingredients'),
    '27066': ('其他医药制品', 'other medical supplies'),
    '35111': ('医疗器械', 'medical equipment'),
}

GROUP_NAMES = {
    '01': ('农林牧渔', 'farming and fishing'), '02': ('煤矿', 'coal mining'), '03': ('石油天然气开采', 'oil and gas'),
    '04': ('金属矿', 'metal ores'), '05': ('非金属矿', 'other mining'), '06': ('食品和烟草', 'food and tobacco'),
    '07': ('纺织', 'textiles'), '08': ('服装鞋帽', 'clothing and shoes'), '09': ('木材家具', 'wood and furniture'),
    '10': ('造纸印刷文体用品', 'paper, printing and stationery'), '11': ('炼油炼焦', 'refining'),
    '12': ('化工', 'chemicals'), '13': ('建材', 'building materials'), '14': ('钢铁有色冶炼', 'metals'),
    '15': ('金属制品', 'metal products'), '16': ('通用设备', 'machinery'), '17': ('专用设备', 'special machinery'),
    '18': ('汽车等交通设备', 'vehicles'), '19': ('电气机械', 'electrical equipment'),
    '20': ('电子和通信设备', 'electronics'), '21': ('仪器仪表', 'instruments'), '22': ('其他制造', 'other manufacturing'),
    '23': ('设备修理', 'equipment repair'), '24': ('电力热力', 'power and heat'), '25': ('燃气', 'gas supply'),
    '26': ('自来水', 'water supply'), '27': ('建筑', 'construction'), '28': ('批发零售', 'wholesale and retail'),
    '29': ('交通运输邮政', 'transport and post'), '30': ('住宿餐饮', 'hotels and restaurants'),
    '31': ('信息技术和互联网', 'IT and internet'), '32': ('金融', 'finance'), '33': ('房地产', 'real estate'),
    '34': ('租赁和商务服务', 'business services'), '35': ('研发', 'research'), '36': ('技术服务', 'technical services'),
    '37': ('公共设施', 'public facilities'), '38': ('居民服务和修理', 'personal services and repair'),
    '39': ('教育', 'education'), '40': ('医疗卫生', 'health care'), '41': ('文化体育娱乐', 'culture and recreation'),
    '42': ('公共管理和社会组织', 'public administration'),
}
DIV_TO_42 = {}
for lo, hi, g in [(1, 5, '01'), (6, 6, '02'), (7, 7, '03'), (8, 9, '04'), (10, 12, '05'), (13, 16, '06'),
                  (17, 17, '07'), (18, 19, '08'), (20, 21, '09'), (22, 24, '10'), (25, 25, '11'), (26, 29, '12'),
                  (30, 30, '13'), (31, 32, '14'), (33, 33, '15'), (34, 34, '16'), (35, 35, '17'), (36, 37, '18'),
                  (38, 38, '19'), (39, 39, '20'), (40, 40, '21'), (41, 42, '22'), (43, 43, '23'), (44, 44, '24'),
                  (45, 45, '25'), (46, 46, '26'), (47, 50, '27'), (51, 52, '28'), (53, 60, '29'), (61, 62, '30'),
                  (63, 65, '31'), (66, 69, '32'), (70, 70, '33'), (71, 72, '34'), (73, 73, '35'), (74, 75, '36'),
                  (76, 79, '37'), (80, 82, '38'), (83, 83, '39'), (84, 85, '40'), (86, 90, '41'), (91, 99, '42')]:
    for d in range(lo, hi + 1):
        DIV_TO_42[d] = g

# id, Chinese label, English label, purchased sector, carries retail margin, bought inside China
ITEMS = [
    ('meal', '一顿饭', 'A restaurant meal', '62172', False, True),
    ('fuel', '一箱油', 'A tank of fuel', '25054', True, False),
    ('phone', '一部手机', 'A phone', '39129', True, False),
    ('rent', '房租', 'Rent', '70183', False, True),
    ('doctor', '看病', 'A doctor visit', '84202', False, True),
    ('school', '学费', 'School fees', '83201', False, True),
    ('clothes', '一件衣服', 'Clothes', '18040', True, False),
    ('liquor', '一瓶酒', 'A bottle of liquor', '15029', True, False),
    ('cigarettes', '一包烟', 'Cigarettes', '16032', True, False),
    ('power', '电费', 'Electricity', '44141', False, True),
]


def num(v):
    return 0.0 if v in (None, '') else float(v)


def load(path):
    digest = hashlib.sha256(open(path, 'rb').read()).hexdigest()
    if digest != EXPECTED_SHA:
        print(f'warning: SHA-256 {digest} differs from the documented file', file=sys.stderr)
    ws = openpyxl.load_workbook(path, data_only=True).worksheets[0]
    assert ws.cell(1, 1).value == '2023年全国投入产出表'
    r0, c0, n = 7, 4, 0
    while ws.cell(r0 + n, 3).value not in ('TII', None):
        n += 1
    codes = [str(ws.cell(r0 + i, 3).value) for i in range(n)]
    names = [ws.cell(r0 + i, 2).value for i in range(n)]
    col = {str(ws.cell(6, c).value): c for c in range(1, ws.max_column + 1) if ws.cell(6, c).value is not None}
    row = {str(ws.cell(r, 3).value): r for r in range(1, ws.max_row + 1) if ws.cell(r, 3).value is not None}
    Z = np.array([[num(ws.cell(r0 + i, c0 + j).value) for j in range(n)] for i in range(n)])
    F = {k: np.array([num(ws.cell(r0 + i, col[k]).value) for i in range(n)]) for k in ['TIU', 'FU102', 'TFU', 'EX', 'IM', 'GO']}
    V = {k: np.array([num(ws.cell(row[k], c0 + j).value) for j in range(n)]) for k in ['VA001', 'VA002', 'VA003', 'VA004']}
    return codes, names, Z, F, V


def main():
    path, out = sys.argv[1], sys.argv[2]
    codes, names, Z, F, V = load(path)
    n = len(codes)
    idx = {c: i for i, c in enumerate(codes)}
    x = F['GO']
    assert np.allclose(Z.sum(0) + sum(V.values()), x, rtol=1e-9, atol=1e-3)
    dom_use = F['TIU'] + F['TFU'] - F['EX']
    s = np.divide(F['IM'], dom_use, out=np.zeros(n), where=dom_use > 0)
    assert (s >= 0).all() and (s <= 1).all()
    A = Z / np.where(x > 0, x, 1)
    Ad = (1 - s)[:, None] * A
    am = (s[:, None] * A).sum(0)
    # wages, profit (operating surplus), depreciation, tax (net production taxes)
    Vc = np.vstack([V['VA001'], V['VA004'], V['VA003'], V['VA002']]) / np.where(x > 0, x, 1)
    L = np.linalg.inv(np.eye(n) - Ad)
    hu = F['FU102']
    retail, whole = idx['52157'], idx['51156']
    goods = [i for i, c in enumerate(codes) if int(c[:2]) <= 43]
    margin = (hu[retail] + hu[whole]) / hu[goods].sum()
    used = set()

    def split(y, domestic_item=None):
        y = y / y.sum() * 100
        share = s.copy()
        if domestic_item is not None:
            share[domestic_item] = 0.0
        direct = float((share * y).sum())
        q = (1 - share) * y
        hops, total = [], np.zeros(5)
        for k in range(HOPS):
            va = Vc @ q
            imp = float(am @ q) + (direct if k == 0 else 0.0)
            nxt = Ad @ q
            order = [i for i in np.argsort(-q) if q[i] >= 0.05][:3]
            for i in order:
                used.add(codes[i])
            row = [round(float(v), 3) for v in va] + [round(imp, 3)]
            hops.append({'keep': row, 'pass': round(float(nxt.sum()), 3),
                         'who': [[codes[i], round(float(q[i]), 2)] for i in order]})
            total += np.array(list(va) + [imp])
            q = nxt
        tail_q = L @ q
        tail = list(Vc @ tail_q) + [float(am @ tail_q)]
        total += np.array(tail)
        whole_q = L @ ((1 - share) * y)
        captured = Vc.sum(0) * whole_q
        groups = {}
        for i, c in enumerate(codes):
            g = DIV_TO_42[int(c[:2])]
            groups[g] = groups.get(g, 0.0) + float(captured[i])
        top_groups = sorted(groups.items(), key=lambda kv: -kv[1])[:6]
        assert abs(total.sum() - 100) < 1e-6, total.sum()
        return {
            'hops': hops,
            'tail': [round(float(v), 3) for v in tail],
            'total': [round(float(v), 3) for v in total],
            'industries': [[g, round(v, 2)] for g, v in top_groups],
        }

    result = {
        'source': 'NBS 2023年全国投入产出表 (211 sectors, competitive-import, producers prices)',
        'order': ['wages', 'profit', 'depreciation', 'tax', 'imports'],
        'retailMargin': round(float(margin), 4),
        'basket': {'id': 'average', **split(hu.copy())},
        'items': [],
    }
    for item_id, zh, en, code, with_margin, domestic in ITEMS:
        y = np.zeros(n)
        i = idx[code]
        y[i] = 1.0
        if with_margin:
            y /= 1 + margin
            y[retail] += margin / (1 + margin) * hu[retail] / (hu[retail] + hu[whole])
            y[whole] += margin / (1 + margin) * hu[whole] / (hu[retail] + hu[whole])
        result['items'].append({'id': item_id, 'zh': zh, 'en': en, 'sector': code,
                                **split(y, i if domestic else None)})
    missing = sorted(c for c in used if c not in SECTOR_NAMES)
    if missing:
        for c in missing:
            print('missing name', c, names[idx[c]], file=sys.stderr)
        if missing:
            sys.exit(1)
    if '--list-sectors' in sys.argv:
        for c in sorted(used):
            print(c, names[idx[c]], '->', SECTOR_NAMES[c][0])
    result['sectors'] = {c: list(SECTOR_NAMES[c]) for c in sorted(used)}
    result['groups'] = {g: list(v) for g, v in GROUP_NAMES.items()}
    with open(out, 'w', encoding='utf-8') as f:
        json.dump(result, f, ensure_ascii=False, separators=(',', ':'))
        f.write('\n')
    b = result['basket']['total']
    print('urban basket per 100:', dict(zip(result['order'], b)), 'margin', result['retailMargin'])


if __name__ == '__main__':
    main()
